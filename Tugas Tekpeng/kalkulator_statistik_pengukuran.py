"""
==================================================================================
 KALKULATOR STATISTIK PENGUKURAN REPLIKASI - VERNIER CALIPER (JANGKA SORONG)
==================================================================================
Tugas   : Teknik Pengukuran - Sistem Pengukuran Replikasi
Objek   : Diameter Baut M16 (diameter nominal 16 mm)
Metode  : 5 orang x 5 kali pengukuran ulang = 25 data (n = 25)

CARA PAKAI (paling sederhana):
    1. Ganti isi dictionary `DATA_PENGUKURAN` di bagian bawah file ini dengan
       25 data hasil pengukuran kelompok kamu (satuan: mm).
    2. Ganti `NILAI_ACUAN` dengan nilai referensi/acuan diameter marker
       (misalnya hasil pengukuran dengan alat yang lebih presisi, atau
       spesifikasi pabrik). Kalau belum ada, boleh dikosongkan (None) --
       maka error & akurasi terhadap nilai acuan tidak akan dihitung.
    3. Ganti `RESOLUSI_ALAT` sesuai skala terkecil vernier caliper kalian
       (umumnya 0.05 mm atau 0.02 mm, cek badan alatnya).
    4. Jalankan:  python kalkulator_statistik_pengukuran.py
    5. File "hasil_statistik_pengukuran.xlsx" akan terbentuk otomatis,
       lengkap dengan rumus Excel (bukan angka mati) sehingga bisa
       diperiksa ulang / diubah datanya langsung di Excel.

Skrip ini murni Python standar (statistics, math) + openpyxl untuk ekspor,
jadi juga siap dipakai sebagai backend untuk web app (Streamlit/Flask) --
lihat contoh `app_streamlit.py` yang disertakan.
==================================================================================
"""

import math
import statistics
from dataclasses import dataclass, field

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference, Series
from openpyxl.chart.label import DataLabelList

# ----------------------------------------------------------------------------
# Tabel nilai kritis distribusi t (dua sisi, 95% CI) untuk beberapa df umum.
# Dipakai untuk menghitung Confidence Interval & Ketidakpastian Pengukuran.
# ----------------------------------------------------------------------------
T_TABLE_95 = {
    1: 12.706, 2: 4.303, 3: 3.182, 4: 2.776, 5: 2.571,
    6: 2.447, 7: 2.365, 8: 2.306, 9: 2.262, 10: 2.228,
    11: 2.201, 12: 2.179, 13: 2.160, 14: 2.145, 15: 2.131,
    16: 2.120, 17: 2.110, 18: 2.101, 19: 2.093, 20: 2.086,
    21: 2.080, 22: 2.074, 23: 2.069, 24: 2.064, 25: 2.060,
    26: 2.056, 27: 2.052, 28: 2.048, 29: 2.045, 30: 2.042,
}


def t_value_95(df: int) -> float:
    if df in T_TABLE_95:
        return T_TABLE_95[df]
    if df > 30:
        return 1.96  # mendekati distribusi normal
    return T_TABLE_95[min(T_TABLE_95.keys(), key=lambda k: abs(k - df))]


@dataclass
class HasilStatistik:
    n: int
    rata_rata: float
    stdev: float
    varians: float
    sem: float                 # standard error of the mean
    median: float
    modus: float
    minimum: float
    maksimum: float
    rentang: float             # range
    cv_persen: float           # coefficient of variation (%)
    presisi_persen: float
    df: int
    t95: float
    ci95_lower: float
    ci95_upper: float
    ketidakpastian: float      # U = t * s / sqrt(n)
    error_mutlak: float = None
    error_relatif_persen: float = None
    akurasi_persen: float = None


def hitung_statistik(data: list, nilai_acuan: float = None) -> HasilStatistik:
    """Menghitung seluruh besaran statistik dari satu himpunan data pengukuran."""
    n = len(data)
    rata = statistics.mean(data)
    sd = statistics.stdev(data)          # sampel (pembagi n-1), sesuai teori ketidakpastian
    var = statistics.variance(data)
    sem = sd / math.sqrt(n)
    med = statistics.median(data)
    try:
        mod = statistics.mode(data)
    except statistics.StatisticsError:
        mod = data[0]
    mn, mx = min(data), max(data)
    rentang = mx - mn
    cv = (sd / rata) * 100 if rata != 0 else float("nan")
    presisi = 100 - cv  # semakin kecil sebaran (CV), semakin tinggi presisi

    df = n - 1
    t95 = t_value_95(df)
    ci_low = rata - t95 * sem
    ci_high = rata + t95 * sem
    ketidakpastian = t95 * sem

    hasil = HasilStatistik(
        n=n, rata_rata=rata, stdev=sd, varians=var, sem=sem, median=med,
        modus=mod, minimum=mn, maksimum=mx, rentang=rentang,
        cv_persen=cv, presisi_persen=presisi, df=df, t95=t95,
        ci95_lower=ci_low, ci95_upper=ci_high, ketidakpastian=ketidakpastian,
    )

    if nilai_acuan is not None:
        hasil.error_mutlak = abs(rata - nilai_acuan)
        hasil.error_relatif_persen = (hasil.error_mutlak / nilai_acuan) * 100
        hasil.akurasi_persen = 100 - hasil.error_relatif_persen

    return hasil


# ==================================================================================
#  EKSPOR KE EXCEL (dengan RUMUS, bukan angka statis, supaya bisa dicek & diubah)
# ==================================================================================

FONT_NAME = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=11)
SUBHEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
SUBHEADER_FONT = Font(name=FONT_NAME, bold=True, size=11)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # kuning = sel yang boleh diisi/diubah user
INPUT_FONT = Font(name=FONT_NAME, color="0000FF")      # biru = data mentah/hardcode input
FORMULA_FONT = Font(name=FONT_NAME, color="000000")    # hitam = hasil rumus
LABEL_FONT = Font(name=FONT_NAME, size=11)
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14)
NOTE_FONT = Font(name=FONT_NAME, italic=True, size=9, color="808080")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header_row(ws, row, col_start, col_end):
    for c in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def buat_workbook(data_per_orang: dict, nilai_acuan, resolusi_alat: float,
                   output_path: str = "hasil_statistik_pengukuran.xlsx"):
    """
    data_per_orang: dict {"Orang 1": [d1..d5], "Orang 2": [...], ...} -> total harus 25 data
    nilai_acuan   : nilai referensi diameter (mm) atau None
    resolusi_alat : skala terkecil vernier caliper yang dipakai (mm)
    """
    orang_list = list(data_per_orang.keys())
    assert len(orang_list) == 5, "Harus ada tepat 5 orang."
    for o in orang_list:
        assert len(data_per_orang[o]) == 5, f"{o} harus punya 5 data pengukuran."

    semua_data = [v for o in orang_list for v in data_per_orang[o]]
    hasil_total = hitung_statistik(semua_data, nilai_acuan)

    wb = openpyxl.Workbook()

    # ------------------------------------------------------------------ SHEET 1: DATA
    ws = wb.active
    ws.title = "Data Pengukuran"
    ws.sheet_view.showGridLines = False

    ws["A1"] = "DATA PENGUKURAN REPLIKASI - DIAMETER BAUT M16 (VERNIER CALIPER)"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    ws["A3"] = "Nilai acuan / referensi diameter (mm):"
    ws["A3"].font = LABEL_FONT
    ws["C3"] = nilai_acuan if nilai_acuan is not None else "-"
    ws["C3"].fill = INPUT_FILL
    ws["C3"].font = INPUT_FONT
    ws["D3"] = "<- isi manual jika ada alat pembanding yang lebih presisi"
    ws["D3"].font = NOTE_FONT

    ws["A4"] = "Resolusi (skala terkecil) alat ukur (mm):"
    ws["A4"].font = LABEL_FONT
    ws["C4"] = resolusi_alat
    ws["C4"].fill = INPUT_FILL
    ws["C4"].font = INPUT_FONT
    ws["D4"] = "<- cek pada rahang vernier caliper (umumnya 0.05 atau 0.02 mm)"
    ws["D4"].font = NOTE_FONT

    header_row = 6
    headers = ["No.", "Nama / Kode Orang", "Ulangan ke-", "Diameter Terukur (mm)"]
    for j, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=j, value=h)
    _style_header_row(ws, header_row, 1, 4)

    row = header_row + 1
    first_data_row = row
    for oi, orang in enumerate(orang_list, start=1):
        for ri, val in enumerate(data_per_orang[orang], start=1):
            ws.cell(row=row, column=1, value=row - first_data_row + 1).border = BORDER
            c2 = ws.cell(row=row, column=2, value=orang)
            c2.border = BORDER
            c3 = ws.cell(row=row, column=3, value=ri)
            c3.border = BORDER
            c4 = ws.cell(row=row, column=4, value=val)
            c4.fill = INPUT_FILL
            c4.font = INPUT_FONT
            c4.border = BORDER
            c4.number_format = "0.00"
            row += 1
    last_data_row = row - 1  # baris terakhir data (harus 25 baris total)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 22

    note_row = last_data_row + 2
    ws.cell(row=note_row, column=1,
             value="Catatan: sel berwarna kuning boleh diganti dengan data pengukuran asli kelompok "
                   "(tetap 5 orang x 5 ulangan = 25 baris). Semua sheet lain otomatis menghitung ulang.")
    ws.cell(row=note_row, column=1).font = NOTE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=4)

    data_range = f"D{first_data_row}:D{last_data_row}"

    # ---------------------------------------------------------- SHEET 2: PER ORANG
    ws2 = wb.create_sheet("Statistik Per Orang")
    ws2.sheet_view.showGridLines = False
    ws2["A1"] = "STATISTIK PER ORANG (REPEATABILITY MASING-MASING PENGUKUR)"
    ws2["A1"].font = TITLE_FONT
    ws2.merge_cells("A1:H1")

    headers2 = ["Orang", "n", "Rata-rata (mm)", "Std. Deviasi (mm)", "Varians",
                "Min (mm)", "Max (mm)", "Rentang (mm)"]
    hr2 = 3
    for j, h in enumerate(headers2, start=1):
        ws2.cell(row=hr2, column=j, value=h)
    _style_header_row(ws2, hr2, 1, len(headers2))

    r = hr2 + 1
    for orang in orang_list:
        rng_diam = f"'Data Pengukuran'!D{first_data_row}:D{last_data_row}"
        rng_nama = f"'Data Pengukuran'!B{first_data_row}:B{last_data_row}"
        ws2.cell(row=r, column=1, value=orang).border = BORDER
        ws2.cell(row=r, column=2,
                  value=f'=COUNTIF({rng_nama},A{r})').border = BORDER
        ws2.cell(row=r, column=3,
                  value=f'=AVERAGEIF({rng_nama},A{r},{rng_diam})').border = BORDER
        ws2.cell(row=r, column=3).number_format = "0.0000"
        # SUMPRODUCT (non-array, aman di semua versi Excel/LibreOffice) untuk STDEV per grup:
        ws2.cell(row=r, column=4,
                  value=(f'=SQRT((SUMPRODUCT(({rng_nama}=A{r})*({rng_diam}-C{r})^2))'
                         f'/(B{r}-1))'))
        ws2.cell(row=r, column=4).number_format = "0.0000"
        ws2.cell(row=r, column=5, value=f'=D{r}^2').border = BORDER
        ws2.cell(row=r, column=5).number_format = "0.000000"
        ws2.cell(row=r, column=6,
                  value=f'=_xlfn.MINIFS({rng_diam},{rng_nama},A{r})').border = BORDER
        ws2.cell(row=r, column=6).number_format = "0.00"
        ws2.cell(row=r, column=7,
                  value=f'=_xlfn.MAXIFS({rng_diam},{rng_nama},A{r})').border = BORDER
        ws2.cell(row=r, column=7).number_format = "0.00"
        ws2.cell(row=r, column=8, value=f'=G{r}-F{r}').border = BORDER
        ws2.cell(row=r, column=8).number_format = "0.00"
        for col in range(1, 9):
            ws2.cell(row=r, column=col).border = BORDER
        r += 1

    for col, w in zip("ABCDEFGH", [16, 6, 16, 18, 12, 12, 12, 14]):
        ws2.column_dimensions[col].width = w

    note2 = r + 1
    ws2.cell(row=note2, column=1,
             value="Rata-rata & SD di sheet ini mengukur 'repeatability' (kekonsistenan) tiap orang. "
                   "Bandingkan dengan sheet 'Statistik Keseluruhan' untuk melihat 'reproducibility' "
                   "(kekonsistenan antar-orang / antar-pengukur).")
    ws2.cell(row=note2, column=1).font = NOTE_FONT
    ws2.merge_cells(start_row=note2, start_column=1, end_row=note2, end_column=8)

    # ---------------------------------------------------------- SHEET 3: RINGKASAN
    ws3 = wb.create_sheet("Statistik Keseluruhan")
    ws3.sheet_view.showGridLines = False
    ws3["A1"] = "RINGKASAN STATISTIK KESELURUHAN (n = 25)"
    ws3["A1"].font = TITLE_FONT
    ws3.merge_cells("A1:C1")

    def add_stat(row, label, formula, fmt="0.0000", note=""):
        ws3.cell(row=row, column=1, value=label).font = LABEL_FONT
        c = ws3.cell(row=row, column=2, value=formula)
        c.number_format = fmt
        c.font = FORMULA_FONT
        c.fill = SUBHEADER_FILL
        c.border = BORDER
        if note:
            ws3.cell(row=row, column=3, value=note).font = NOTE_FONT

    dref = f"'Data Pengukuran'!{data_range}"
    n_cell = f"COUNT({dref})"

    rows_def = []
    rows_def.append(("Jumlah data (n)", f"={n_cell}", "0", "Total pengukuran (5 orang x 5 ulangan)"))
    rows_def.append(("Rata-rata (mean), mm", f"=AVERAGE({dref})", "0.0000", ""))
    rows_def.append(("Standar Deviasi (s), mm", f"=STDEV({dref})", "0.0000",
                      "Sebaran data terhadap rata-rata (n-1)"))
    rows_def.append(("Varians (s^2)", f"=VAR({dref})", "0.000000", ""))
    rows_def.append(("Standard Error of Mean (SEM), mm", f"=B{{r_sd}}/SQRT(B{{r_n}})", "0.0000",
                      "Ketelitian estimasi rata-rata"))
    rows_def.append(("Median, mm", f"=MEDIAN({dref})", "0.0000", ""))
    rows_def.append(("Modus, mm", f"=IFERROR(MODE({dref}),\"tidak ada\")", "0.0000", ""))
    rows_def.append(("Nilai Minimum, mm", f"=MIN({dref})", "0.00", ""))
    rows_def.append(("Nilai Maksimum, mm", f"=MAX({dref})", "0.00", ""))
    rows_def.append(("Rentang (Range), mm", f"=B{{r_max}}-B{{r_min}}", "0.00", ""))
    rows_def.append(("Koefisien Variasi (CV), %", f"=(B{{r_sd}}/B{{r_mean}})*100", "0.00", "Ukuran presisi relatif"))
    rows_def.append(("Presisi, %", f"=100-B{{r_cv}}", "0.00", "Semakin mendekati 100%, semakin presisi"))
    rows_def.append(("Derajat kebebasan (df = n-1)", f"=B{{r_n}}-1", "0", ""))
    rows_def.append(("Nilai t (95% CI, dua sisi)", f"={hasil_total.t95}", "0.000",
                      "Dari tabel distribusi-t, df=" + str(hasil_total.df)))
    rows_def.append(("Batas bawah CI 95%, mm", f"=B{{r_mean}}-B{{r_t}}*B{{r_sem}}", "0.0000", ""))
    rows_def.append(("Batas atas CI 95%, mm", f"=B{{r_mean}}+B{{r_t}}*B{{r_sem}}", "0.0000", ""))
    rows_def.append(("Ketidakpastian Pengukuran (U = t.SEM), mm", f"=B{{r_t}}*B{{r_sem}}", "0.0000",
                      "Laporkan hasil sebagai: rata-rata +/- U"))
    rows_def.append(("Resolusi alat ukur, mm", f"='Data Pengukuran'!C4", "0.00", ""))
    rows_def.append(("Ketidakpastian alat (1/2 resolusi), mm", f"=B{{r_res}}/2", "0.0000",
                      "Ketidakpastian minimal akibat skala terkecil alat"))

    rows_def.append(("UCL - Upper Control Limit (mean + 3s), mm", f"=B{{r_mean}}+3*B{{r_sd}}", "0.0000",
                      "Batas kontrol atas (peta kendali 3-sigma)"))
    rows_def.append(("LCL - Lower Control Limit (mean - 3s), mm", f"=B{{r_mean}}-3*B{{r_sd}}", "0.0000",
                      "Batas kontrol bawah (peta kendali 3-sigma)"))

    if nilai_acuan is not None:
        rows_def.append(("Nilai Acuan (True Value), mm", "='Data Pengukuran'!C3", "0.0000", ""))
        rows_def.append(("Error Mutlak (Absolute Error), mm", f"=ABS(B{{r_mean}}-B{{r_ref}})", "0.0000", ""))
        rows_def.append(("Error Relatif, %", f"=(B{{r_err}}/B{{r_ref}})*100", "0.00", ""))
        rows_def.append(("Akurasi, %", f"=100-B{{r_erel}}", "0.00", "Semakin mendekati 100%, semakin akurat"))

    start_row = 3
    refs = {}
    for i, (label, formula, fmt, note) in enumerate(rows_def):
        r = start_row + i
        refs[label] = r

    def R(label):
        return refs[label]

    for i, (label, formula, fmt, note) in enumerate(rows_def):
        r = start_row + i
        formula = formula.format(
            r_sd=R("Standar Deviasi (s), mm"),
            r_n=R("Jumlah data (n)"),
            r_mean=R("Rata-rata (mean), mm"),
            r_max=R("Nilai Maksimum, mm"),
            r_min=R("Nilai Minimum, mm"),
            r_cv=R("Koefisien Variasi (CV), %"),
            r_t=R("Nilai t (95% CI, dua sisi)"),
            r_sem=R("Standard Error of Mean (SEM), mm"),
            r_res=R("Resolusi alat ukur, mm"),
            r_ref=R("Nilai Acuan (True Value), mm") if nilai_acuan is not None else "",
            r_err=R("Error Mutlak (Absolute Error), mm") if nilai_acuan is not None else "",
            r_erel=R("Error Relatif, %") if nilai_acuan is not None else "",
        ) if "{" in formula else formula
        add_stat(r, label, formula, fmt, note)

    ws3.column_dimensions["A"].width = 38
    ws3.column_dimensions["B"].width = 14
    ws3.column_dimensions["C"].width = 46

    kesimpulan_row = start_row + len(rows_def) + 2
    ws3.cell(row=kesimpulan_row, column=1, value="Cara membaca hasil:").font = SUBHEADER_FONT
    ws3.cell(row=kesimpulan_row + 1, column=1,
             value="- SD kecil & CV rendah -> presisi tinggi (data antar ulangan konsisten).")
    ws3.cell(row=kesimpulan_row + 2, column=1,
             value="- Error relatif kecil terhadap nilai acuan -> akurasi tinggi (dekat dengan nilai sebenarnya).")
    ws3.cell(row=kesimpulan_row + 3, column=1,
             value="- Data bisa presisi tapi tidak akurat (konsisten tapi meleset dari nilai acuan), atau sebaliknya.")
    for rr in range(kesimpulan_row + 1, kesimpulan_row + 4):
        ws3.cell(row=rr, column=1).font = NOTE_FONT
        ws3.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)

    # ---------------------------------------------------------- SHEET 4: GRAFIK
    _tambah_sheet_grafik(wb, orang_list, first_data_row, refs, nilai_acuan)

    wb.save(output_path)
    return output_path, hasil_total


def _tambah_sheet_grafik(wb, orang_list, first_data_row, refs_ringkasan, nilai_acuan):
    """
    Membuat sheet 'Grafik' berisi tabel bantu (link formula ke sheet Data) dan
    chart NATIVE Excel (LineChart/BarChart). Chart ini interaktif di Excel:
    bisa di-hover untuk lihat nilai per titik, dan otomatis ter-update kalau
    data pengukuran diubah.
    """
    ws = wb.create_sheet("Grafik")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "GRAFIK INTERAKTIF (otomatis mengikuti perubahan data)"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:H1")

    n_orang = len(orang_list)
    R = lambda label: refs_ringkasan[label]

    # ================= Tabel bantu 1: data per ulangan per orang (untuk line chart) ====
    t1_row = 3
    ws.cell(row=t1_row, column=1, value="Tabel Bantu 1 - Diameter per Ulangan per Orang").font = SUBHEADER_FONT
    hdr = t1_row + 1
    ws.cell(row=hdr, column=1, value="Ulangan")
    for j, orang in enumerate(orang_list, start=2):
        ws.cell(row=hdr, column=j, value=orang)
    _style_header_row(ws, hdr, 1, n_orang + 1)
    for u in range(1, 6):
        r = hdr + u
        ws.cell(row=r, column=1, value=u).border = BORDER
        for p, orang in enumerate(orang_list):
            src_row = first_data_row + p * 5 + (u - 1)
            c = ws.cell(row=r, column=p + 2, value=f"='Data Pengukuran'!D{src_row}")
            c.number_format = "0.0000"
            c.border = BORDER
    t1_last = hdr + 5

    line1 = LineChart()
    line1.title = "Hasil Tiap Ulangan per Orang (Repeatability)"
    line1.x_axis.title = "Ulangan ke-"
    line1.y_axis.title = "Diameter (mm)"
    line1.style = 12
    cats = Reference(ws, min_col=1, min_row=hdr + 1, max_row=t1_last)
    data = Reference(ws, min_col=2, max_col=n_orang + 1, min_row=hdr, max_row=t1_last)
    line1.add_data(data, titles_from_data=True)
    line1.set_categories(cats)
    for s in line1.series:
        s.marker.symbol = "circle"
        s.smooth = False
    line1.height, line1.width = 9, 16
    ws.add_chart(line1, f"A{t1_last + 2}")

    # ================= Tabel bantu 2: rata-rata & SD per orang (ambil dari sheet lain) ==
    t2_row = t1_last + 20
    ws.cell(row=t2_row, column=1, value="Tabel Bantu 2 - Rata-rata & SD per Orang").font = SUBHEADER_FONT
    hdr2 = t2_row + 1
    for j, h in enumerate(["Orang", "Rata-rata (mm)", "SD (mm)"], start=1):
        ws.cell(row=hdr2, column=j, value=h)
    _style_header_row(ws, hdr2, 1, 3)
    for i, orang in enumerate(orang_list):
        r = hdr2 + 1 + i
        src = 4 + i  # baris data di sheet 'Statistik Per Orang' (header di baris 3)
        ws.cell(row=r, column=1, value=f"='Statistik Per Orang'!A{src}").border = BORDER
        ws.cell(row=r, column=2, value=f"='Statistik Per Orang'!C{src}").number_format = "0.0000"
        ws.cell(row=r, column=2).border = BORDER
        ws.cell(row=r, column=3, value=f"='Statistik Per Orang'!D{src}").number_format = "0.0000"
        ws.cell(row=r, column=3).border = BORDER
    t2_last = hdr2 + n_orang

    bar2 = BarChart()
    bar2.type = "col"
    bar2.title = "Rata-rata per Orang (Reproducibility Antar-Pengukur)"
    bar2.y_axis.title = "Diameter (mm)"
    bar2.style = 10
    cats2 = Reference(ws, min_col=1, min_row=hdr2 + 1, max_row=t2_last)
    data2 = Reference(ws, min_col=2, min_row=hdr2, max_row=t2_last)
    bar2.add_data(data2, titles_from_data=True)
    bar2.set_categories(cats2)
    bar2.height, bar2.width = 9, 14
    ws.add_chart(bar2, f"E{t1_last + 2}")

    # ================= Tabel bantu 3: histogram (8 bin) =================================
    t3_row = t2_last + 3
    ws.cell(row=t3_row, column=1, value="Tabel Bantu 3 - Histogram (8 kelas)").font = SUBHEADER_FONT
    hdr3 = t3_row + 1
    for j, h in enumerate(["Kelas", "Batas Bawah", "Batas Atas", "Frekuensi"], start=1):
        ws.cell(row=hdr3, column=j, value=h)
    _style_header_row(ws, hdr3, 1, 4)
    dref = "'Data Pengukuran'!D{}:D{}".format(first_data_row, first_data_row + 5 * n_orang - 1)
    n_bin = 8
    for k in range(n_bin):
        r = hdr3 + 1 + k
        ws.cell(row=r, column=1, value=f"Kelas {k+1}").border = BORDER
        ws.cell(row=r, column=2,
                 value=(f"='Statistik Keseluruhan'!B{R('Nilai Minimum, mm')}"
                        f"+({k}/{n_bin})*('Statistik Keseluruhan'!B{R('Nilai Maksimum, mm')}"
                        f"-'Statistik Keseluruhan'!B{R('Nilai Minimum, mm')})")).number_format = "0.0000"
        ws.cell(row=r, column=2).border = BORDER
        ws.cell(row=r, column=3,
                 value=(f"='Statistik Keseluruhan'!B{R('Nilai Minimum, mm')}"
                        f"+({k+1}/{n_bin})*('Statistik Keseluruhan'!B{R('Nilai Maksimum, mm')}"
                        f"-'Statistik Keseluruhan'!B{R('Nilai Minimum, mm')})")).number_format = "0.0000"
        ws.cell(row=r, column=3).border = BORDER
        last_bin = "1" if k == n_bin - 1 else "0"
        ws.cell(row=r, column=4,
                 value=(f'=COUNTIFS({dref},">="&B{r},{dref},"<"&C{r})'
                        + (f'+COUNTIF({dref},"="&C{r})' if k == n_bin - 1 else ""))
                 ).border = BORDER
    t3_last = hdr3 + n_bin

    bar3 = BarChart()
    bar3.type = "col"
    bar3.title = "Distribusi (Histogram) Seluruh 25 Data"
    bar3.y_axis.title = "Frekuensi"
    bar3.x_axis.title = "Kelas Diameter"
    bar3.style = 11
    cats3 = Reference(ws, min_col=1, min_row=hdr3 + 1, max_row=t3_last)
    data3 = Reference(ws, min_col=4, min_row=hdr3, max_row=t3_last)
    bar3.add_data(data3, titles_from_data=True)
    bar3.set_categories(cats3)
    bar3.gapWidth = 10
    bar3.height, bar3.width = 9, 14
    ws.add_chart(bar3, f"A{t1_last + 22}")

    # ================= Tabel bantu 4: control chart (25 data + CL/UCL/LCL) ==============
    t4_row = t3_last + 3
    ws.cell(row=t4_row, column=1, value="Tabel Bantu 4 - Control Chart (25 data berurutan)").font = SUBHEADER_FONT
    hdr4 = t4_row + 1
    for j, h in enumerate(["No.", "Diameter (mm)", "CL", "UCL", "LCL"], start=1):
        ws.cell(row=hdr4, column=j, value=h)
    _style_header_row(ws, hdr4, 1, 5)
    total_n = 5 * n_orang
    for i in range(total_n):
        r = hdr4 + 1 + i
        src_row = first_data_row + i
        ws.cell(row=r, column=1, value=i + 1).border = BORDER
        ws.cell(row=r, column=2, value=f"='Data Pengukuran'!D{src_row}").number_format = "0.0000"
        ws.cell(row=r, column=2).border = BORDER
        ws.cell(row=r, column=3, value=f"='Statistik Keseluruhan'!B{R('Rata-rata (mean), mm')}").number_format = "0.0000"
        ws.cell(row=r, column=3).border = BORDER
        ws.cell(row=r, column=4,
                 value=f"='Statistik Keseluruhan'!B{R('UCL - Upper Control Limit (mean + 3s), mm')}").number_format = "0.0000"
        ws.cell(row=r, column=4).border = BORDER
        ws.cell(row=r, column=5,
                 value=f"='Statistik Keseluruhan'!B{R('LCL - Lower Control Limit (mean - 3s), mm')}").number_format = "0.0000"
        ws.cell(row=r, column=5).border = BORDER
    t4_last = hdr4 + total_n

    line4 = LineChart()
    line4.title = "Control Chart - 25 Data Berurutan (dengan CL/UCL/LCL)"
    line4.x_axis.title = "Urutan Pengukuran"
    line4.y_axis.title = "Diameter (mm)"
    line4.style = 13
    cats4 = Reference(ws, min_col=1, min_row=hdr4 + 1, max_row=t4_last)
    data4 = Reference(ws, min_col=2, max_col=5, min_row=hdr4, max_row=t4_last)
    line4.add_data(data4, titles_from_data=True)
    line4.set_categories(cats4)
    for s, dashed in zip(line4.series, [False, True, True, True]):
        s.marker.symbol = "circle" if not dashed else "none"
        s.smooth = False
    line4.height, line4.width = 9, 20
    ws.add_chart(line4, f"A{t1_last + 42}")

    # ================= Tabel bantu 5: Presisi vs Akurasi (gauge) ========================
    t5_row = t4_last + 3
    ws.cell(row=t5_row, column=1, value="Tabel Bantu 5 - Presisi vs Akurasi (%)").font = SUBHEADER_FONT
    hdr5 = t5_row + 1
    ws.cell(row=hdr5, column=1, value="Besaran")
    ws.cell(row=hdr5, column=2, value="Nilai (%)")
    _style_header_row(ws, hdr5, 1, 2)
    ws.cell(row=hdr5 + 1, column=1, value="Presisi").border = BORDER
    ws.cell(row=hdr5 + 1, column=2,
             value=f"='Statistik Keseluruhan'!B{R('Presisi, %')}").number_format = "0.00"
    ws.cell(row=hdr5 + 1, column=2).border = BORDER
    t5_last = hdr5 + 1
    if nilai_acuan is not None:
        ws.cell(row=hdr5 + 2, column=1, value="Akurasi").border = BORDER
        ws.cell(row=hdr5 + 2, column=2,
                 value=f"='Statistik Keseluruhan'!B{R('Akurasi, %')}").number_format = "0.00"
        ws.cell(row=hdr5 + 2, column=2).border = BORDER
        t5_last = hdr5 + 2

    bar5 = BarChart()
    bar5.type = "bar"
    bar5.title = "Ringkasan Presisi vs Akurasi Keseluruhan"
    bar5.x_axis.title = "%"
    bar5.style = 15
    cats5 = Reference(ws, min_col=1, min_row=hdr5 + 1, max_row=t5_last)
    data5 = Reference(ws, min_col=2, min_row=hdr5, max_row=t5_last)
    bar5.add_data(data5, titles_from_data=True)
    bar5.set_categories(cats5)
    bar5.height, bar5.width = 6, 12
    ws.add_chart(bar5, f"E{t1_last + 42}")

    # ================= Tabel bantu 6: Error & Akurasi per orang ==========================
    if nilai_acuan is not None:
        t6_row = t5_last + 3
        ws.cell(row=t6_row, column=1, value="Tabel Bantu 6 - Error & Akurasi per Orang").font = SUBHEADER_FONT
        hdr6 = t6_row + 1
        for j, h in enumerate(["Orang", "Error Mutlak (mm)", "Akurasi (%)"], start=1):
            ws.cell(row=hdr6, column=j, value=h)
        _style_header_row(ws, hdr6, 1, 3)
        for i, orang in enumerate(orang_list):
            r = hdr6 + 1 + i
            gref_row = hdr2 + 1 + i  # baris rata-rata di Tabel Bantu 2
            ws.cell(row=r, column=1, value=orang).border = BORDER
            ws.cell(row=r, column=2,
                     value=f"=ABS(B{gref_row}-'Data Pengukuran'!C3)").number_format = "0.0000"
            ws.cell(row=r, column=2).border = BORDER
            ws.cell(row=r, column=3,
                     value=f"=100-(B{r}/'Data Pengukuran'!C3*100)").number_format = "0.00"
            ws.cell(row=r, column=3).border = BORDER
        t6_last = hdr6 + n_orang

        bar6 = BarChart()
        bar6.type = "col"
        bar6.title = "Akurasi per Orang (%)"
        bar6.y_axis.title = "Akurasi (%)"
        bar6.style = 14
        cats6 = Reference(ws, min_col=1, min_row=hdr6 + 1, max_row=t6_last)
        data6 = Reference(ws, min_col=3, min_row=hdr6, max_row=t6_last)
        bar6.add_data(data6, titles_from_data=True)
        bar6.set_categories(cats6)
        bar6.height, bar6.width = 9, 14
        ws.add_chart(bar6, f"A{t1_last + 62}")
    else:
        t6_last = t5_last

    # ================= Tabel bantu 7: CV / Presisi per orang =============================
    t7_row = t6_last + 3
    ws.cell(row=t7_row, column=1, value="Tabel Bantu 7 - Presisi (CV) per Orang").font = SUBHEADER_FONT
    hdr7 = t7_row + 1
    for j, h in enumerate(["Orang", "CV (%)", "Presisi (%)"], start=1):
        ws.cell(row=hdr7, column=j, value=h)
    _style_header_row(ws, hdr7, 1, 3)
    for i, orang in enumerate(orang_list):
        r = hdr7 + 1 + i
        mean_ref = hdr2 + 1 + i
        ws.cell(row=r, column=1, value=orang).border = BORDER
        ws.cell(row=r, column=2,
                 value=f"=(C{mean_ref}/B{mean_ref})*100").number_format = "0.00"
        ws.cell(row=r, column=2).border = BORDER
        ws.cell(row=r, column=3, value=f"=100-B{r}").number_format = "0.00"
        ws.cell(row=r, column=3).border = BORDER
    t7_last = hdr7 + n_orang

    bar7 = BarChart()
    bar7.type = "col"
    bar7.title = "Presisi (100% - CV) per Orang"
    bar7.y_axis.title = "Presisi (%)"
    bar7.style = 16
    cats7 = Reference(ws, min_col=1, min_row=hdr7 + 1, max_row=t7_last)
    data7 = Reference(ws, min_col=3, min_row=hdr7, max_row=t7_last)
    bar7.add_data(data7, titles_from_data=True)
    bar7.set_categories(cats7)
    bar7.height, bar7.width = 9, 14
    ws.add_chart(bar7, f"E{t1_last + 62}")

    # ================= Tabel bantu 8: Ringkasan Batas Kontrol ============================
    t8_row = t7_last + 3
    ws.cell(row=t8_row, column=1, value="Tabel Bantu 8 - Ringkasan Batas Kontrol").font = SUBHEADER_FONT
    hdr8 = t8_row + 1
    ws.cell(row=hdr8, column=1, value="Batas")
    ws.cell(row=hdr8, column=2, value="Nilai (mm)")
    _style_header_row(ws, hdr8, 1, 2)
    label_refs = [
        ("LCL (-3s)", "LCL - Lower Control Limit (mean - 3s), mm"),
        ("CI 95% bawah", "Batas bawah CI 95%, mm"),
        ("CL (rata-rata)", "Rata-rata (mean), mm"),
        ("CI 95% atas", "Batas atas CI 95%, mm"),
        ("UCL (+3s)", "UCL - Upper Control Limit (mean + 3s), mm"),
    ]
    for i, (label, key) in enumerate(label_refs):
        r = hdr8 + 1 + i
        ws.cell(row=r, column=1, value=label).border = BORDER
        ws.cell(row=r, column=2, value=f"='Statistik Keseluruhan'!B{R(key)}").number_format = "0.0000"
        ws.cell(row=r, column=2).border = BORDER
    t8_last = hdr8 + len(label_refs)

    bar8 = BarChart()
    bar8.type = "col"
    bar8.title = "Batas Kontrol & Confidence Interval"
    bar8.y_axis.title = "Diameter (mm)"
    bar8.style = 17
    cats8 = Reference(ws, min_col=1, min_row=hdr8 + 1, max_row=t8_last)
    data8 = Reference(ws, min_col=2, min_row=hdr8, max_row=t8_last)
    bar8.add_data(data8, titles_from_data=True)
    bar8.set_categories(cats8)
    bar8.height, bar8.width = 9, 14
    ws.add_chart(bar8, f"A{t1_last + 82}")

    for col, w in zip("ABCDEFGH", [22, 16, 14, 14, 14, 14, 14, 14]):
        ws.column_dimensions[col].width = w


# ==================================================================================
#  CONTOH PEMAKAIAN / TEMPLATE DATA -- GANTI DENGAN DATA ASLI KELOMPOK KAMU
# ==================================================================================
if __name__ == "__main__":
    # Data 25 pengukuran diameter baut M16 (mm) - dari data_pengukuran_baut_M16.xlsx
    DATA_PENGUKURAN = {
        "Orang 1": [15.98, 16.00, 15.98, 16.02, 16.00],
        "Orang 2": [16.04, 16.02, 16.06, 16.04, 16.02],
        "Orang 3": [15.92, 15.94, 15.92, 15.96, 15.94],
        "Orang 4": [16.00, 15.98, 16.02, 16.00, 15.98],
        "Orang 5": [15.96, 15.98, 15.96, 16.00, 15.98],
    }

    NILAI_ACUAN = 16.00       # mm - diameter nominal baut M16 (bisa diganti None jika tidak dipakai)
    RESOLUSI_ALAT = 0.02      # mm - skala terkecil vernier caliper yang dipakai

    path, hasil = buat_workbook(DATA_PENGUKURAN, NILAI_ACUAN, RESOLUSI_ALAT)

    print(f"File berhasil dibuat: {path}")
    print("-" * 50)
    print(f"n              : {hasil.n}")
    print(f"Rata-rata      : {hasil.rata_rata:.4f} mm")
    print(f"Std. Deviasi   : {hasil.stdev:.4f} mm")
    print(f"SEM            : {hasil.sem:.4f} mm")
    print(f"CV             : {hasil.cv_persen:.2f} %")
    print(f"Presisi        : {hasil.presisi_persen:.2f} %")
    print(f"Ketidakpastian : +/- {hasil.ketidakpastian:.4f} mm (95% CI)")
    if hasil.error_mutlak is not None:
        print(f"Error Mutlak   : {hasil.error_mutlak:.4f} mm")
        print(f"Error Relatif  : {hasil.error_relatif_persen:.2f} %")
        print(f"Akurasi        : {hasil.akurasi_persen:.2f} %")
